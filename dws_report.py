# -*- coding: utf-8 -*-
"""สรุปเวลาลงพัสดุแยกตามช่อง DWS

หลักการ
    คลังสแกน (dws_collect) บอกว่า "ใบงานนี้ไปลงช่องไหน กี่ชิ้น"
    Report เส้นหลัก/เส้นรอง บอกว่า "ใบงานนี้ใช้เวลาลงพัสดุเท่าไหร่"
    join ด้วยหมายเลขใบงาน แล้วรวมเวลา group by ช่อง

ทำไมต้องเก็บ Report สะสมด้วย
    คอลัมน์เวลาลงพัสดุจะกรอกก็ต่อเมื่อรถลงของเสร็จแล้ว ตอนบอทดึง Report รอบแรก
    ของ window นั้นค่ายังว่างอยู่ ต้องดึงซ้ำในรอบหลัง ๆ ถึงจะได้ค่า จึง upsert
    ทับลงคลังทุกรอบ ค่าที่มีเวลาแล้วจะไม่ถูกค่าว่างเขียนทับ

หน่วยเวลา
    เส้นหลัก  "ระยะเวลาลงพัสดุ (ชั่วโมง)"   = ชั่วโมง
    เส้นรอง   "ระยะเวลาขนถ่ายลงรถ(min）"    = นาที → หาร 60

ใบงานที่ลงมากกว่า 1 ช่อง
    หารเวลาตามสัดส่วนจำนวนชิ้นที่สแกนได้ในแต่ละช่อง
"""

import os
import sqlite3

import pandas as pd

from dws_collect import _connect, load_scans, store_path

# (คำที่ต้องมีครบทุกคำในชื่อคอลัมน์, ตัวคูณให้เป็นชั่วโมง)
DURATION_COLUMNS = [
    (("ระยะเวลาลงพัสดุ",), 1.0),        # เส้นหลัก AY — ชั่วโมงอยู่แล้ว
    (("ขนถ่ายลงรถ", "min"), 1.0 / 60),  # เส้นรอง BA — นาที
]

# กันไปจับคอลัมน์ "เวลารอคิว" ที่ชื่อใกล้กันมาก
# เส้นรองมีทั้ง "รายะเวลาเข้าแถวขนถ่ายลงรถ(min）" และ "ระยะเวลาขนถ่ายลงรถ(min）"
EXCLUDE_KEYWORDS = ("เข้าแถว", "เข้าคิว", "รอลง")

START_KEYWORDS = ("เริ่มงานที่ใช้ขนย้ายลงรถ",)
END_KEYWORDS = ("จบงานที่ใช้ขนย้ายลงรถ",)


def _find_column(df, keyword_sets):
    for keywords in keyword_sets:
        for col in df.columns:
            text = str(col)
            if any(bad in text for bad in EXCLUDE_KEYWORDS):
                continue
            if all(k in text for k in keywords):
                return col
    return None


def _find_duration_column(df):
    for keywords, factor in DURATION_COLUMNS:
        col = _find_column(df, [keywords])
        if col is not None:
            return col, factor
    return None, None


def hhmm(hours):
    """ทศนิยมชั่วโมง → ข้อความ ชม:นาที เช่น 0.55 → "0:33" และ 19.72 → "19:43"

    ทศนิยมชั่วโมงอ่านยากเวลาเทียบกันด้วยตา (0.55 กับ 33 นาที ไม่ได้เชื่อมกันในหัวทันที)
    """
    if hours is None or hours != hours:
        return ""
    total = int(round(float(hours) * 60))
    return f"{total // 60}:{total % 60:02d}"


def _ensure_tasks_table(con):
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS tasks (
            task    TEXT PRIMARY KEY,
            line    TEXT,
            hours   REAL,
            started TEXT,
            ended   TEXT
        )
        """
    )
    con.commit()


def upsert_tasks(store_folder, report_path, line_name, log=print):
    """อ่าน Report แล้ว upsert เวลาลงพัสดุลงคลัง คืนจำนวนใบที่มีเวลา"""
    report_path = (report_path or "").strip()
    if not report_path or not os.path.exists(report_path):
        log(f"  {line_name}: ไม่พบไฟล์ {report_path or '(ไม่ได้ระบุ)'}")
        return 0

    df = pd.read_excel(report_path, dtype=str).fillna("")
    col, factor = _find_duration_column(df)
    if col is None:
        log(f"  {line_name}: ไม่พบคอลัมน์เวลาลงพัสดุ")
        return 0

    start_col = _find_column(df, [START_KEYWORDS])
    end_col = _find_column(df, [END_KEYWORDS])

    rows = pd.DataFrame({
        "task": df.iloc[:, 0].str.strip(),
        "line": line_name,
        "hours": pd.to_numeric(df[col], errors="coerce") * factor,
        "started": df[start_col].str.strip() if start_col else "",
        "ended": df[end_col].str.strip() if end_col else "",
    })
    rows = rows[(rows["task"] != "") & rows["hours"].notna()]

    con = _connect(store_path(store_folder))
    try:
        _ensure_tasks_table(con)
        # ค่าที่มีเวลาแล้วเท่านั้นที่ถูกเขียนลงไป ค่าว่างจึงไม่ทับของเดิม
        con.executemany(
            """
            INSERT INTO tasks (task, line, hours, started, ended) VALUES (?,?,?,?,?)
            ON CONFLICT(task) DO UPDATE SET
                line=excluded.line, hours=excluded.hours,
                started=excluded.started, ended=excluded.ended
            """,
            rows[["task", "line", "hours", "started", "ended"]].itertuples(index=False, name=None),
        )
        con.commit()
        total = con.execute("SELECT COUNT(*) FROM tasks").fetchone()[0]
    finally:
        con.close()

    log(f"  {line_name}: {len(df)} ใบ | มีเวลา {len(rows)} ใบ "
        f"({len(rows) / max(len(df), 1) * 100:.0f}%) | คลังสะสม {total:,} ใบ")
    return len(rows)


def load_tasks(store_folder):
    db = store_path(store_folder)
    if not os.path.exists(db):
        return pd.DataFrame(columns=["task", "line", "hours"])
    con = _connect(db)
    try:
        _ensure_tasks_table(con)
        return pd.read_sql_query("SELECT task, line, hours FROM tasks", con)
    finally:
        con.close()


def build_summary(store_folder, start_time=None, end_time=None, split="scans"):
    """คืน (สรุปตามช่อง, รายใบงาน, ใบที่ยังไม่มีเวลา)"""
    scans = load_scans(store_folder, start_time, end_time)
    tasks = load_tasks(store_folder)

    detail = scans.merge(tasks, on="task", how="left")
    if detail.empty:
        empty = pd.DataFrame()
        return pd.DataFrame(index=pd.Index(range(1, 12), name="ช่อง DWS")), empty, empty

    total_scans = detail.groupby("task")["scans"].transform("sum")
    bay_count = detail.groupby("task")["bay"].transform("size")
    if split == "equal":
        detail["share"] = 1.0 / bay_count
    elif split == "max":
        top = detail.groupby("task")["scans"].transform("max")
        detail["share"] = (detail["scans"] == top).astype(float)
        detail["share"] /= detail.groupby("task")["share"].transform("sum")
    elif split == "full":
        detail["share"] = 1.0
    else:  # scans — ตามสัดส่วนจำนวนชิ้น
        detail["share"] = detail["scans"] / total_scans

    detail["hours_bay"] = detail["hours"] * detail["share"]
    detail["multi"] = bay_count > 1

    matched = detail.dropna(subset=["hours"])
    missing = detail[detail["hours"].isna()][["task", "bay", "scans"]]

    summary = (
        matched.groupby("bay")
        .agg(ใบงาน=("task", "nunique"), เวลารวม=("hours_bay", "sum"), สแกน=("scans", "sum"))
        .reindex(range(1, 12))
    )
    summary["ใบงาน"] = summary["ใบงาน"].astype("Int64")
    summary["สแกน"] = summary["สแกน"].astype("Int64")
    # "ชิ้นต่อชั่วโมง" คิดจากชั่วโมงเต็มก่อนปัด ไม่งั้นความคลาดเคลื่อนจากการปัด
    # จะถูกคูณต่อและสะสมขึ้นเมื่อมีใบงานเยอะ
    summary["ชิ้นต่อชั่วโมง"] = (summary["สแกน"] / summary["เวลารวม"]).round(0)
    summary = summary[["ใบงาน", "เวลารวม", "สแกน", "ชิ้นต่อชั่วโมง"]]
    summary.index.name = "ช่อง DWS"

    detail = detail.rename(columns={
        "task": "หมายเลขใบงาน", "bay": "ช่อง", "scans": "จำนวนสแกน",
        "line": "เส้นทาง", "hours": "เวลาลงพัสดุ(ชม.)",
        "share": "สัดส่วนที่แบ่งให้ช่องนี้", "hours_bay": "เวลาที่นับให้ช่องนี้(ชม.)",
        "multi": "ลงหลายช่อง",
    })
    # ใส่คอลัมน์อ่านง่ายคู่กับทศนิยมไว้ตรวจย้อนกลับ
    detail.insert(
        detail.columns.get_loc("เวลาที่นับให้ช่องนี้(ชม.)") + 1,
        "เวลาที่นับให้ช่องนี้ (ชม:นาที)",
        detail["เวลาที่นับให้ช่องนี้(ชม.)"].map(hhmm),
    )
    missing = missing.rename(columns={
        "task": "หมายเลขใบงาน", "bay": "ช่อง", "scans": "จำนวนสแกน",
    })
    return summary, detail, missing


def _style_summary_sheet(ws, n_rows, meta_start=None):
    """แต่งชีตสรุปให้ถ่ายเป็นรูปส่ง Feishu แล้วอ่านออก"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Tahoma", size=11)
    idle_fill = PatternFill("solid", fgColor="F2F2F2")

    for cell in ws[1]:
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 32

    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1):
        idle = row[1].value is None          # ช่องที่ไม่ได้ใช้งาน
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
            if idle:
                cell.fill = idle_fill
        if not idle:
            row[3].number_format = "#,##0"   # สแกน

    # คอลัมน์ C = เวลารวม เก็บเป็นค่าเวลาจริงของ Excel (เศษของวัน) แล้วใส่
    # format [h]:mm — แสดงเป็น 0:33 / 19:43 แต่ยัง sum และ sort ได้ตามปกติ
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1, min_col=3, max_col=3):
        cell = row[0]
        if isinstance(cell.value, (int, float)):
            cell.value = float(cell.value) / 24.0
            cell.number_format = "[h]:mm"

    for col, width in zip("ABCDE", (11, 9, 13, 11, 15)):
        ws.column_dimensions[col].width = width

    if meta_start:
        for row in ws.iter_rows(min_row=meta_start, max_row=ws.max_row):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(horizontal="left")


def write_excel(output_path, summary, detail, missing, meta=None):
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    meta_start = None
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="สรุปตามช่อง")
        if meta:
            meta_start = len(summary) + 4
            pd.DataFrame(list(meta.items()), columns=["หัวข้อ", "ค่า"]).to_excel(
                writer, sheet_name="สรุปตามช่อง", index=False,
                startrow=meta_start - 1,
            )
        _style_summary_sheet(writer.sheets["สรุปตามช่อง"], len(summary), meta_start)
        if len(detail):
            detail.to_excel(writer, sheet_name="รายใบงาน", index=False)
        if len(missing):
            missing.to_excel(writer, sheet_name="ยังไม่มีเวลา", index=False)
    return output_path
